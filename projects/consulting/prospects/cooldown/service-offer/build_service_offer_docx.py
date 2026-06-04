from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "service-offer/Annabel_Filippini_Small_Business_Website_AI_Services.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="DADCE0", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)


def add_package_heading(doc, name, price=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(name)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if price:
        price_run = paragraph.add_run(f" - {price}")
        price_run.bold = True
        price_run.font.size = Pt(12)
        price_run.font.color.rgb = RGBColor(85, 85, 85)


def add_includes(doc, items):
    for item in items:
        add_bullet(doc, item)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15

for style_name, size, before, after in [
    ("Title", 24, 0, 10),
    ("Heading 1", 16, 18, 8),
    ("Heading 2", 14, 14, 6),
    ("Heading 3", 12, 10, 4),
]:
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

title = doc.add_paragraph(style="Title")
title.add_run("Practical AI Support + Website Redesign For Small Businesses")

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
subtitle_run = subtitle.add_run("Start with the website. Keep going into AI, marketing, and automation only where it creates real value.")
subtitle_run.italic = True
subtitle_run.font.color.rgb = RGBColor(85, 85, 85)

doc.add_paragraph(
    "Hi, I'm Annabel. I help small businesses use AI in practical, non-overwhelming ways, starting with the place customers already interact with you most: your website."
)
doc.add_paragraph(
    "The first project can be a website redesign or refresh. If that is all you want, we can stop there. If it opens up bigger opportunities, like marketing systems, content workflows, customer follow-up, or admin automation, we can keep going from there."
)
doc.add_paragraph(
    "My best fit is a founder, owner, or small team that already has a business people love, but has a website, marketing process, or internal workflow that feels harder than it needs to be."
)

doc.add_heading("Why AI Matters For Small Businesses", level=1)
doc.add_paragraph(
    "AI is becoming a real advantage for small teams because it helps with the work that usually gets delayed: writing, organizing, following up, researching, repurposing content, and keeping marketing consistent."
)
doc.add_paragraph("A few current signals:")
for item in [
    "QuickBooks reported in 2025 that 68% of small businesses surveyed now use AI regularly, up from 48% the previous summer, and 74% said AI is boosting productivity.",
    "Thryv's 2025 small business survey found AI usage increased from 39% in 2024 to 55% in 2025, with content marketing called out as a common use case.",
    "Paychex reported that among small businesses using AI, 66% saw increased productivity, 44% saw cost savings, 40% saw revenue growth, and 34% saw improved customer acquisition.",
    "McKinsey estimates generative AI can improve marketing productivity by 5% to 15% of total marketing spend and customer-care productivity by 30% to 45% of function costs.",
]:
    add_bullet(doc, item)
doc.add_paragraph(
    "The opportunity is not to use AI everywhere. The opportunity is to find the few places where AI can make your business feel lighter, faster, and more consistent."
)

doc.add_heading("What I Can Help With", level=1)
doc.add_heading("Website Redesign and Cleanup", level=2)
doc.add_paragraph("I can help refresh or rebuild the parts of your website that matter most:")
for item in [
    "Homepage messaging and layout",
    "Product or service pages",
    "Navigation and mobile experience",
    "Calls to action, contact forms, booking links, and purchase flows",
    "Shopify theme cleanup and owner-friendly editing setup",
    "Basic SEO improvements like page titles, meta descriptions, alt text, and cleaner page structure",
    "A simple handoff guide so your team knows how to make routine updates afterward",
]:
    add_bullet(doc, item)

doc.add_paragraph(
    "The goal is not just to make the site prettier. The goal is to make it easier for customers to understand what you offer, trust the business, and take the next step."
)

doc.add_heading("Recent Example: Cool Down", level=2)
doc.add_paragraph(
    "I recently worked on a Shopify redesign for Cool Down, a running and apparel brand. The work focused on making the site easier to shop, easier to edit, and more aligned with how the business is growing."
)
doc.add_paragraph("Examples of the work included:")
for item in [
    "Improving product page clarity so customers could better understand size, color, availability, and product details",
    "Cleaning up theme structure so routine edits could happen inside Shopify instead of custom code",
    "Improving run club and location page structure",
    "Creating a simple editing guide for the team",
    "Identifying quick technical fixes, like broken promotional links and missing SEO/accessibility basics",
]:
    add_bullet(doc, item)

doc.add_paragraph(
    "That project is the type of work I would love to bring to other small businesses: practical, visible improvements that make the site work better for both customers and the owner."
)

doc.add_heading("AI And Automation Support", level=1)
doc.add_paragraph(
    "I can help with AI systems when there is a clear workflow worth improving. This is optional, and I usually recommend starting with the website first because it shows where your customer journey, messaging, and operations are already getting stuck."
)
for item in [
    "Drafting email outreach from a list of leads",
    "Turning notes, events, or photos into newsletter/social copy",
    "Creating reusable prompts or custom AI workspaces for your team's voice",
    "Organizing customer, sponsor, or partner research",
    "Automating repetitive admin steps between tools like forms, email, docs, spreadsheets, and calendars",
    "Creating simple content systems so you can turn one idea, event, product launch, or customer story into multiple useful marketing pieces",
]:
    add_bullet(doc, item)
doc.add_paragraph(
    "My approach is simple: I look for the repetitive task that is taking up too much time, then build the smallest useful system that saves time without making the business more complicated."
)

doc.add_heading("Marketing And Content Systems", level=1)
doc.add_paragraph(
    "I can also help you look at what you are already doing for marketing and where AI could make it easier to stay consistent."
)
doc.add_paragraph("That might mean:")
for item in [
    "Clarifying what your website should say before someone buys, books, visits, or reaches out",
    "Turning product launches, events, customer stories, or behind-the-scenes moments into repeatable content prompts",
    "Creating a weekly newsletter or social content workflow",
    "Building lightweight outreach systems for partners, local businesses, creators, sponsors, or customers",
    "Helping you evaluate whether your current marketing feels comfortable, sustainable, and aligned with your voice",
]:
    add_bullet(doc, item)
doc.add_paragraph(
    "One strategy you may have heard of is an engagement pod or influencer pod, where creators coordinate in a group chat to like, comment on, and share each other's posts quickly. I understand why people use them, but I would usually recommend a more authentic version for small businesses: a creator seeding, ambassador, or community content sprint."
)
doc.add_paragraph(
    "That means identifying people who genuinely like the business, giving them a clear reason to participate, making it easy for them to post, and using AI to help with prompts, follow-up, tracking, and content repurposing. The goal is not fake engagement. The goal is more real people creating more real proof around the business."
)

doc.add_heading("Ways To Work Together", level=1)
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Inches(1.55)
table.columns[1].width = Inches(1.15)
table.columns[2].width = Inches(1.2)
table.columns[3].width = Inches(2.6)
headers = ["Package", "Payment", "Timeline", "Best Fit"]
for idx, header in enumerate(headers):
    cell = table.rows[0].cells[idx]
    cell.text = header
    set_cell_shading(cell, "F1F3F4")
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True
rows = [
    ["Website Refresh Sprint", "Choose your price", "1 to 2 weeks", "Existing site that needs polish, cleanup, and clearer customer flow."],
    ["Website Redesign Sprint", "Choose your price", "2 to 4 weeks", "Most small businesses: a meaningful redesign without a full agency engagement."],
    ["Website + AI Operations Sprint", "Choose your price", "3 to 5 weeks", "A stronger website plus one practical AI or marketing workflow."],
    ["Ongoing Support", "Discuss together", "Monthly", "Continued site updates, launches, cleanup, and AI workflow support."],
]
for row in rows:
    cells = table.add_row().cells
    for idx, value in enumerate(row):
        cells[idx].text = value
        set_cell_margins(cells[idx])
        cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
set_table_borders(table)

add_package_heading(doc, "Website Refresh Sprint")
doc.add_paragraph("Best for a business with an existing site that needs polish, cleanup, and clearer customer flow.")
add_includes(doc, [
    "Site review and prioritized punch list",
    "Refresh of 1 to 3 key pages",
    "Mobile layout cleanup",
    "Navigation and call-to-action improvements",
    "Basic SEO/accessibility cleanup",
    "30-minute handoff call",
    "Simple update guide",
])

add_package_heading(doc, "Website Redesign Sprint")
doc.add_paragraph("Best for a business that wants the site to feel meaningfully redesigned without hiring a full agency.")
add_includes(doc, [
    "Site strategy and structure review",
    "Redesign of up to 5 core pages or templates",
    "Shopify or website editor setup so the owner can make routine updates",
    "Product/service page improvements",
    "Mobile-first layout refinement",
    "Basic SEO/accessibility cleanup",
    "One revision round",
    "Handoff guide and walkthrough video",
])
doc.add_paragraph("This is the package I would recommend for most small businesses Bailey sends my way.")

add_package_heading(doc, "Website + AI Operations Sprint")
doc.add_paragraph("Best for a small team that wants both a stronger website and one practical AI or marketing workflow.")
add_includes(doc, [
    "Everything in the Website Redesign Sprint",
    "One AI system such as outreach drafting, content assistance, creator seeding support, inquiry response support, research organization, or a lightweight admin workflow",
])

add_package_heading(doc, "Ongoing Support")
doc.add_paragraph("For businesses that want continued help after the initial sprint.")
add_includes(doc, [
    "Monthly website updates",
    "New product or service launches",
    "Small design fixes",
    "SEO/content cleanup",
    "AI workflow refinement",
    "Owner/team support as new needs come up",
])

doc.add_heading("Choose-Your-Price Approach", level=1)
doc.add_paragraph(
    "Because AI is changing how website work gets done, I am currently offering this as a choose-your-price engagement for referred small businesses."
)
doc.add_paragraph("Here is how it works:")
for item in [
    "I review your website and recommend the smallest useful scope.",
    "We agree on what I will deliver and what timeline makes sense.",
    "You choose a price that feels fair based on the value to your business, your current budget, and the amount of work involved.",
]:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run(item)

doc.add_paragraph(
    "I care most about doing useful work, learning where this creates the most value, and building relationships with businesses I genuinely want to help."
)
doc.add_paragraph(
    "If it is helpful to have a reference point, small business website refreshes like this often land somewhere between $1,000 and $5,000+, depending on scope. That is not a fixed rate. It is just a guide so you have context when choosing what feels fair."
)
doc.add_paragraph(
    "For larger rebuilds, urgent timelines, or ongoing support, we can talk openly about what would make the work sustainable for both sides."
)

doc.add_heading("How We Would Start", level=1)
for item in [
    "You send me your website and a quick note about what feels frustrating or outdated.",
    "I do a short review of the site, customer journey, and any marketing/workflow pain points you want me to consider.",
    "I send back the biggest opportunities I see and recommend the smallest useful first project.",
    "We start with the website if that is the clearest win, then only expand into AI automation or marketing systems if it feels valuable.",
]:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run(item)

doc.add_paragraph(
    "I like working with businesses that are already doing good work and need their website and systems to finally keep up."
)
doc.add_paragraph(
    "If Bailey sent you my name, I would be happy to take a look and tell you honestly where I think I could help."
)

signature = doc.add_paragraph()
signature.alignment = WD_ALIGN_PARAGRAPH.LEFT
signature_run = signature.add_run("Annabel Filippini")
signature_run.bold = True

doc.add_heading("Sources", level=1)
for item in [
    "QuickBooks Small Business Insights Survey, 2025: https://quickbooks.intuit.com/r/small-business-data/april-2025-survey/",
    "Thryv AI and Small Business Survey, 2025: https://www.businesswire.com/news/home/20250717239434/en/AI-Adoption-Among-Small-Businesses-Surges-41-in-2025-According-to-New-Survey-from-Thryv",
    "Paychex AI Small Business Survey, 2025: https://www.nasdaq.com/press-release/paychex-survey-finds-ai-empowering-small-businesses-72-feeling-positive-2025-03-18",
    "McKinsey, The economic potential of generative AI: https://www.mckinsey.com/capabilities/mckinsey-digital/our-insights/the-economic-potential-of-generative-ai-the-next-productivity-frontier",
    "Hootsuite definition of engagement pods: https://blog.hootsuite.com/social-media-definitions/engagement-pod/",
    "Sprout Social guide to influencer seeding: https://sproutsocial.com/insights/influencer-seeding/",
]:
    add_bullet(doc, item)

doc.save(OUT)
print(OUT)
